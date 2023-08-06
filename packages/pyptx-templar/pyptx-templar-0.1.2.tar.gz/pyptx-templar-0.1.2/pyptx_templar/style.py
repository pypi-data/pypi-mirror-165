from .xml import SubElement
from pptx.dml.color import RGBColor as RGB
from pptx.enum.dml import MSO_LINE
from pptx.util import Pt


# La dernière bordure ajoutée est gardée (à l'ajout d'une bordure, l'ancienne
# est retirée si elle existe).
# border_color est la couleur de la bordure
# border_width est la largeur de la bordure
# border_style est le style de la bordure
# left, top, right, bot indiquent quelles bordures sont à afficher
def set_cell_borders(cell,
                     border_color=RGB(0, 0, 0),
                     border_width=Pt(1),
                     border_style=MSO_LINE.DASH,
                     left=False,
                     top=False,
                     right=False,
                     bot=False):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    borders = [('L', left), ('R', right), ('T', top), ('B', bot)]
    for idx, (side, do) in enumerate(borders):
        if not do:
            continue
        tag = f"a:ln{side}"

        # On retire la bordure précédente (si elle existe).
        tcPr.remove_all(tag)

        # les bordures doivent être déclarées dans un certain ordre (L, R, T, B)
        # sans quoi elles ne sont pas affichées...
        before = [f"a:ln{s}" for (s, _) in borders[idx + 1:]]
        # les bordures doivent apparaitre dans le xml avant les style de fond de la cellule
        # sinon elles ne sont pas affichées...
        before += ['a:solidFill', 'a:noFill']

        ln = SubElement(tcPr,
                        tag,
                        before=before,
                        w=str(border_width),
                        cap='flat',
                        cmpd='sng',
                        algn='ctr')

        solidFill = SubElement(ln, 'a:solidFill')
        # couleur rgb:
        colorval = "".join(map(lambda v: f"{v:02X}", border_color))
        SubElement(solidFill, 'a:srgbClr', val=colorval)
        # couleur prédéfinie:
        # SubElement(solidFill, 'a:schemeClr', val="tx1")

        SubElement(ln, 'a:prstDash', val=MSO_LINE.to_xml(border_style))
        SubElement(ln, 'a:round')
        SubElement(ln, 'a:headEnd', type='none', w='med', len='med')
        SubElement(ln, 'a:tailEnd', type='none', w='med', len='med')


# left, top, right, bot sont les coordonnées des cellules aux bordures
# outer indique si l'on veut afficher les bordures extérieures
# inner indique si l'on veut afficher les bordures intérieures
def set_rect_borders(table,
                     left,
                     top,
                     right,
                     bot,
                     outer=False,
                     inner=False,
                     **kwargs):
    assert left >= 0
    assert top >= 0
    assert right < len(table.columns)
    assert bot < len(table.rows)
    assert left <= right
    assert top <= bot

    for x in range(left, right + 1):
        for y in range(top, bot + 1):
            topb = outer if y == top else inner
            botb = outer if y == bot else inner
            leftb = outer if x == left else inner
            rightb = outer if x == right else inner

            set_cell_borders(table.cell(y, x),
                             **kwargs,
                             left=leftb,
                             top=topb,
                             right=rightb,
                             bot=botb)

    if not outer:
        return

    # PowerPoint a un bug (?), pour les cases intérieures du tableau
    # il ne regarde que les bordures droite et bas, il faut donc aussi
    # ajouter les bordures sur les cases extérieures de la selection

    for y in range(top, bot + 1):
        # bordure droite de la colonne de gauche
        if left != 0:
            set_cell_borders(table.cell(y, left - 1), **kwargs, right=True)
        # bordure gauche de la colonne de droite
        if right != len(table.columns) - 1:
            set_cell_borders(table.cell(y, right + 1), **kwargs, left=True)

    for x in range(left, right + 1):
        # bordure inférieure de la ligne du haut
        if top != 0:
            set_cell_borders(table.cell(top - 1, x), **kwargs, bot=True)
        # bordure supérieure de la ligne du bas
        if bot != len(table.rows) - 1:
            set_cell_borders(table.cell(bot + 1, x), **kwargs, top=True)


# size est la taille de la police, en général en Pt ou Cm:
# https://python-pptx.readthedocs.io/en/latest/api/util.html
# font est le nom de la font à utiliser (chaine de caractères, sensible à la casse)
# bold indique si le texte doit être en gras
# italic indique si le texte doit être en italique
# font_rgb est la couleur du texte, de type RGBColor, construit soit à partir de 3 entiers
# soit à partir d'une chaine de caractères représentant 3 hexadécimaux:
# https://python-pptx.readthedocs.io/en/latest/api/dml.html#pptx.dml.color.RGBColor
# font_brightness est la luminosité de la police, il s'agit d'un flottant entre -1 et 1
# underline indique si le texte doit être souligné
# language indique la langue du texte:
# https://python-pptx.readthedocs.io/en/latest/api/enum/MsoLanguageId.html
def font_style(ft,
               size=None,
               font=None,
               bold=None,
               italic=None,
               font_rgb=None,
               font_brightness=None,
               underline=None,
               language=None):
    if size is not None:
        ft.size = size
    if font is not None:
        ft.name = font
    if bold is not None:
        ft.bold = bold
    if italic is not None:
        ft.italic = italic
    if font_rgb is not None:
        ft.color.rgb = font_rgb
    if font_brightness is not None:
        ft.color.brightness = font_brightness
    if underline is not None:
        ft.underline = underline
    if language is not None:
        ft.language_id = language


# halign indique l'alignement horizontal du texte:
# https://python-pptx.readthedocs.io/en/latest/api/enum/PpParagraphAlignment.html
def text_style(tf, halign=None, **kwargs):
    for paragraph in tf.paragraphs:
        font_style(paragraph.font, **kwargs)
        if halign is not None:
            paragraph.alignment = halign


# valign indique l'alignement vertical du texte
# https://python-pptx.readthedocs.io/en/latest/api/enum/MsoVerticalAnchor.html
# back_rgb est la couleur de fond de la cellule
# https://python-pptx.readthedocs.io/en/latest/api/dml.html#pptx.dml.color.RGBColor
# back_brightness est la luminosité du fond de la cellule, un float entre -1 et 1
# margin_top, margin_left, margin_bottom, margin_right sont les marges de la cellule
# https://python-pptx.readthedocs.io/en/latest/api/util.html
def cell_style(cell,
               valign=None,
               back_rgb=None,
               back_brightness=None,
               margin_bottom=None,
               margin_top=None,
               margin_left=None,
               margin_right=None,
               **kwargs):
    if margin_bottom is not None:
        cell.margin_bottom = margin_bottom
    if margin_top is not None:
        cell.margin_top = margin_top
    if margin_right is not None:
        cell.margin_right = margin_right
    if margin_left is not None:
        cell.margin_left = margin_left
    if valign is not None:
        cell.vertical_anchor = valign
    if back_brightness is not None:
        cell.fill.solid()
        cell.fill.fore_color.brightness = back_brightness
    if back_rgb is not None:
        if not back_rgb:
            cell.fill.background()
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = back_rgb

    text_style(cell.text_frame, **kwargs)


def paragraph_append(paragraph, txt, **kwargs):
    run = paragraph.add_run()
    run.text = txt
    font_style(run.font, **kwargs)
