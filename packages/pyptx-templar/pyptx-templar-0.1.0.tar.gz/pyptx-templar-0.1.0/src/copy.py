import logging
import sys

from pptx.oxml.text import CT_RegularTextRun, CT_TextCharacterProperties
from pptx.oxml.text import CT_TextLineBreak, CT_TextParagraphProperties

from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL


def color_copy(cfrom, cto):
    if cfrom.type is not None:
        if cfrom.type == MSO_COLOR_TYPE.RGB:
            cto.rgb = cfrom.rgb
        elif cfrom.type == MSO_COLOR_TYPE.SCHEME:
            cto.theme_color = cfrom.theme_color
        else:
            logging.warning("font_copy: type de couleur inconnu '%s'",
                            cfrom.type)
        cto.brightness = cfrom.brightness


def fill_copy(ffrom, fto):
    if ffrom.type is None:
        return

    if ffrom.type == MSO_FILL.BACKGROUND:
        fto.background()
    elif ffrom.type == MSO_FILL.GRADIENT:
        fto.gradient()
        fto.gradient_angle = ffrom.gradient_angle
        fto.gradient_stops = ffrom.gradient_stops
    elif ffrom.type == MSO_FILL.PATTERNED:
        fto.patterned()
        fto.pattern = ffrom.pattern
        color_copy(ffrom.back_color, fto.back_color)
    elif ffrom.type == MSO_FILL.PICTURE:
        fto.picture()
    elif ffrom.type == MSO_FILL.SOLID:
        fto.solid()
    elif ffrom.type == MSO_FILL.TEXTURED:
        fto.textured()
    else:
        logging.warning("font_copy: type de remplissage inconnu '%s'",
                        ffrom.type)
    color_copy(ffrom.fore_color, fto.fore_color)


def font_copy(ffrom, fto):
    fto.bold = ffrom.bold
    fto.italic = ffrom.italic
    fto.language_id = ffrom.language_id
    fto.name = ffrom.name
    fto.size = ffrom.size
    fto.underline = ffrom.underline
    fill_copy(ffrom.fill, fto.fill)

    # attention accéder au champ color met le fill type à SOLID
    # il ne faut donc le faire que si ffrom.type == SOLID
    # (sinon ffrom.type est changé à SOLID)
    if ffrom.fill.type == MSO_FILL.SOLID:
        color_copy(ffrom.color, fto.color)


def run_copy(rfrom, rto):
    font_copy(rfrom.font, rto.font)
    rto.text = rfrom.text


def paragraph_copy(pfrom, pto):
    pto.clear()

    run_idx = 0
    # on itère directement sur les éléments xml car
    # c'est la seule manière de récupérer les sauts de ligne
    for e in pfrom._element:
        if isinstance(e, CT_RegularTextRun):
            # du texte (un run)
            run = pto.add_run()
            run_copy(pfrom.runs[run_idx], run)
            run_idx += 1
        elif isinstance(e, CT_TextLineBreak):
            # un saut de ligne
            pto.add_line_break()
        elif not isinstance(
                e, (CT_TextParagraphProperties, CT_TextCharacterProperties)):
            # CT_TextParagraphProperties est au début de chaque paragraphe et
            # contient les information de style du paragraphe, déjà copiées plus bas
            # CT_TextCharacterProperties est au début des runs qui ont un style différent,
            # et contient les informations sur ce style, copiées plus haut par run_copy
            logging.warning("paragraph_copy: element non prévu de type '%s'",
                            type(e))

    # il faut les copier après avoir créé les runs sans quoi
    # le style n'est pas complètement gardé
    # (add_line_break casse le niveau ?)
    font_copy(pfrom.font, pto.font)
    pto.alignment = pfrom.alignment
    pto.line_spacing = pfrom.line_spacing
    pto.space_after = pfrom.space_after
    pto.space_before = pfrom.space_before
    pto.level = pfrom.level


def textframe_copy(tffrom, tfto):
    tfto.clear()

    while len(tfto.paragraphs) != len(tffrom.paragraphs):
        pt = tfto.add_paragraph()

    for pf, pt in zip(tffrom.paragraphs, tfto.paragraphs):
        paragraph_copy(pf, pt)

    tfto.margin_bottom = tffrom.margin_bottom
    tfto.margin_left = tffrom.margin_left
    tfto.margin_right = tffrom.margin_right
    tfto.margin_top = tffrom.margin_top
    tfto.vertical_anchor = tffrom.vertical_anchor
    tfto.word_wrap = tffrom.word_wrap
